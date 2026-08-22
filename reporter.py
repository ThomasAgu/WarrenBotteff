import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class TradeReporter:
    def __init__(self, reports_dir="reports"):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        self.reports_dir = os.path.join(base_dir, reports_dir)
        os.makedirs(self.reports_dir, exist_ok=True)

    def _get_filepath(self, date_str=None):
        date_str = date_str or datetime.now().strftime("%Y-%m-%d")
        return os.path.join(self.reports_dir, f"Operaciones_{date_str}.xlsx")

    def _style_sheet(self, ws, money_columns=(), percent_columns=()):
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Segoe UI", size=10)
        border = Border(*(Side(style="thin", color="E0E0E0") for _ in range(4)))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = border
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F9FBFD")
            for column in money_columns:
                ws.cell(row=row[0].row, column=column).number_format = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'
            for column in percent_columns:
                ws.cell(row=row[0].row, column=column).number_format = '0.00%;[Red](0.00%);"-"'
        for column_cells in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 3
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 14)
        ws.freeze_panes = "A2"

    def _write_trades(self, wb, trades):
        ws = wb.create_sheet("Trades")
        ws.append(["Entry Date", "Entry Time", "Exit Date", "Exit Time", "Symbol", "Entry Price", "Exit Price", "Quantity", "Gross PnL", "Fees", "Net PnL", "Return %", "Duration (seconds)"])
        for trade in trades:
            ws.append([trade.entry_time.date(), trade.entry_time.time(), trade.exit_time.date(), trade.exit_time.time(), trade.symbol, trade.entry_price, trade.exit_price, trade.quantity, trade.gross_pnl, trade.fees, trade.net_pnl, trade.return_percentage / 100, trade.duration_seconds])
        self._style_sheet(ws, money_columns=(6, 7, 9, 10, 11), percent_columns=(12,))

    def _write_equity(self, wb, equity_curve):
        ws = wb.create_sheet("Equity")
        ws.append(["Timestamp", "Equity", "Drawdown", "Drawdown %"])
        peak = None
        for timestamp, equity in sorted(equity_curve, key=lambda point: point[0]):
            peak = equity if peak is None else max(peak, equity)
            drawdown = equity - peak
            ws.append([timestamp, equity, drawdown, drawdown / peak if peak else 0.0])
        self._style_sheet(ws, money_columns=(2, 3), percent_columns=(4,))

    def _write_summary(self, wb, metrics, strategy_name):
        ws = wb.create_sheet("Summary")
        ws.append(["Metric", "Value"])
        ws.append(["Strategy", strategy_name])
        labels = [("Initial Capital", "initial_balance"), ("Final Capital", "final_balance"), ("Net Profit", "net_profit"), ("Return %", "return_percentage"), ("Total Trades", "total_trades"), ("Winning Trades", "winning_trades"), ("Losing Trades", "losing_trades"), ("Win Rate", "win_rate"), ("Gross Profit", "gross_profit"), ("Gross Loss", "gross_loss"), ("Profit Factor", "profit_factor"), ("Total Fees", "total_fees"), ("Average Trade", "average_trade"), ("Average Win", "average_win"), ("Average Loss", "average_loss"), ("Average Trade Duration", "average_trade_duration_seconds"), ("Max Drawdown", "max_drawdown"), ("Max Drawdown %", "max_drawdown_percentage"), ("Sharpe Ratio", "sharpe_ratio"), ("Sortino Ratio", "sortino_ratio"), ("Start Time", "start_time"), ("End Time", "end_time"), ("Duration (seconds)", "duration_seconds")]
        for label, key in labels:
            value = metrics[key]
            if key in ("return_percentage", "win_rate", "max_drawdown_percentage"):
                value /= 100
            ws.append([label, value])
        self._style_sheet(ws, money_columns=(2,))
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value in ("Return %", "Win Rate", "Max Drawdown %"):
                ws.cell(row=row, column=2).number_format = '0.00%;[Red](0.00%);"-"'

    def generate_report(self, trades, metrics, equity_curve, strategy_name="EMA_RSI", date_str=None):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self._write_trades(wb, trades)
        self._write_equity(wb, equity_curve)
        self._write_summary(wb, metrics, strategy_name)
        filepath = self._get_filepath(date_str)
        wb.save(filepath)
        return filepath
